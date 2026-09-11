from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from src.common.data_loader import file_hash
from src.q2.result_writer import emergency_intervals, slot_intervals, write_outputs

ROOT = Path(__file__).resolve().parents[1]


class ResultWriterTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.template = ROOT / "data/raw/附件5/result2.xlsx"
        cls.plan = pd.read_csv(ROOT / "outputs/q2/q2_plan_schedule.csv", float_precision="round_trip")
        cls.actual = pd.read_csv(ROOT / "outputs/q2/q2_actual_schedule.csv", float_precision="round_trip")
        cls.daily = pd.read_csv(ROOT / "outputs/q2/q2_daily_metrics.csv", float_precision="round_trip")
        workbook = load_workbook(cls.template)
        cls.labels = [workbook["计划购电量"].cell(1, c).value for c in range(2, 146)]
        cls.intervals = slot_intervals(cls.labels)
        workbook.close()

    def test_emergency_grouping_and_no_cross_day(self):
        frame = self.actual.iloc[:288].copy()
        frame["emergency_purchase_kwh"] = 0.0
        frame.loc[[0, 1, 3, 143, 144], "emergency_purchase_kwh"] = 2.0
        frame.loc[2, "emergency_purchase_kwh"] = 1e-8
        events = emergency_intervals(frame, self.intervals)
        self.assertEqual(len(events), 4)
        self.assertEqual(events[0]["interval"], "0:10-0:30")
        self.assertEqual(events[0]["emergency_purchase_kwh"], 4.)
        self.assertEqual(events[2]["interval"], "0:00+1-0:10+1")
        self.assertNotEqual(events[2]["date"], events[3]["date"])
        self.assertEqual(sum(r["emergency_purchase_kwh"] for r in events), 10.)

    def test_expanded_workbook_and_immutable_csvs(self):
        frozen = [self.template] + list((ROOT / "outputs/q2").glob("*.csv")) + [ROOT / "outputs/schedule/q1_schedule.csv"]
        before = {p: file_hash(p) for p in frozen}
        source = load_workbook(self.template)
        try:
            with TemporaryDirectory() as directory:
                path, paper = Path(directory) / "result2.xlsx", Path(directory) / "tables.md"
                events = write_outputs(self.template, path, self.plan, self.actual, self.daily, paper)
                result = load_workbook(path)
                try:
                    self.assertEqual(result.sheetnames, source.sheetnames)
                    planned, battery, emergency = [result[s] for s in result.sheetnames]
                    self.assertEqual((planned.max_row, planned.max_column), (335, 147))
                    values = np.array([[planned.cell(r, c).value for c in range(2, 146)] for r in range(2, 336)])
                    np.testing.assert_allclose(values.ravel(), self.plan.grid_purchase_plan_kwh, rtol=0, atol=1e-8)
                    self.assertEqual(battery.max_row, 2005)
                    self.assertEqual(str(battery.merged_cells), str(source["充放电量"].merged_cells))
                    for day in range(334):
                        first = 2 + day * 6
                        actual = self.actual.iloc[day * 144:(day + 1) * 144]
                        self.assertEqual(battery.cell(first, 1).value.date().isoformat(), actual.date.iloc[0])
                        self.assertAlmostEqual(battery.cell(first, 6).value, actual.soc_real_start_kwh.iloc[0], delta=1e-8)
                        self.assertAlmostEqual(battery.cell(first + 1, 6).value, actual.soc_real_end_kwh.iloc[-1], delta=1e-8)
                        for block in range(6):
                            part = actual.iloc[block * 24:(block + 1) * 24]
                            self.assertAlmostEqual(battery.cell(first + block, 3).value, (part.x_real_kwh + part.q_real_kwh).sum(), delta=1e-7)
                            self.assertAlmostEqual(battery.cell(first + block, 4).value, part.z_real_kwh.sum(), delta=1e-7)
                    for row in (2, 3, 7, 1999, 2005):
                        source_row = 2 + (row - 2) % 6
                        self.assertEqual(battery.row_dimensions[row].height, source["充放电量"].row_dimensions[source_row].height)
                        for col in range(1, 7):
                            self.assertEqual(battery.cell(row, col)._style or result._cell_styles[0],
                                             source["充放电量"].cell(source_row, col)._style or source._cell_styles[0])
                    self.assertEqual(emergency.max_row, len(events) + 1)
                    for row, event in enumerate(events, 2):
                        self.assertEqual(emergency.cell(row, 1).value.date().isoformat(), event["date"])
                        self.assertEqual(emergency.cell(row, 2).value, event["interval"])
                        self.assertAlmostEqual(emergency.cell(row, 3).value, event["emergency_purchase_kwh"], delta=1e-8)
                    self.assertAlmostEqual(sum(e["emergency_purchase_kwh"] for e in events),
                                           self.actual.loc[self.actual.emergency_purchase_kwh > 1e-6, "emergency_purchase_kwh"].sum(), delta=1e-7)
                    for name in source.sheetnames:
                        self.assertEqual({k: dict(v) for k, v in source[name].column_dimensions.items()},
                                         {k: dict(v) for k, v in result[name].column_dimensions.items()})
                    self.assertIn("2025-03-20", paper.read_text(encoding="utf-8"))
                    self.assertIn("2025-12-21", paper.read_text(encoding="utf-8"))
                finally:
                    result.close()
        finally:
            source.close()
        self.assertEqual(before, {p: file_hash(p) for p in frozen})

    def test_missing_date_rejected(self):
        with TemporaryDirectory() as directory:
            with self.assertRaisesRegex(ValueError, "不完整"):
                write_outputs(self.template, Path(directory) / "bad.xlsx", self.plan.iloc[:-1], self.actual,
                              self.daily, Path(directory) / "bad.md")
