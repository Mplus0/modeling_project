"""最终固定参数、验收门槛、结果比较和官方模板映射的轻量测试。"""

import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from src.common.data_loader import file_hash
from src.q3.final import FINAL_ALPHA, FINAL_LAMBDA, FINAL_LABEL, read_confirmed_selection, compare_search_winner, validate_final, load_final
from src.q3.result_writer import inspect_template, _populate, verify_workbook, write_result3
from src.q3.final_reports import comparison_table, write_paper_reports
from src.q3.annual import FORMAL_DATES

ROOT = Path(__file__).resolve().parents[1]


class FinalTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.ranking,cls.selection,cls.daily,cls.metrics = read_confirmed_selection(ROOT)

    def test_fixed_parameters_and_cli_override_rejected(self):
        self.assertEqual((FINAL_ALPHA,FINAL_LAMBDA),(.85,.25))
        for option in ("--alpha","--lambda"):
            result = subprocess.run([sys.executable,str(ROOT/"scripts/10_run_q3_final.py"),option,"0.90"],capture_output=True,text=True)
            self.assertEqual(result.returncode,2)
            self.assertIn("unrecognized arguments",result.stderr)

    def test_winner_search_identity_and_rejection(self):
        final = self.daily.assign(run_label=FINAL_LABEL)
        metrics = dict(self.metrics,run_label=FINAL_LABEL,runtime_seconds=1.)
        self.assertTrue(compare_search_winner(final,metrics,self.daily,self.metrics)["passed"])
        final.loc[0,"total_actual_cost_yuan"] += .01
        self.assertFalse(compare_search_winner(final,metrics,self.daily,self.metrics)["passed"])
        self.assertFalse(compare_search_winner(self.daily.iloc[:-1],metrics,self.daily,self.metrics)["passed"])
        altered = dict(metrics,total_emergency_cost_yuan=0)
        self.assertFalse(compare_search_winner(self.daily,altered,self.daily,self.metrics)["passed"])

    def test_q2_comparison_recalculation(self):
        q2 = json.loads((ROOT/"outputs/q2/q2_metrics.json").read_text(encoding="utf-8"))
        comparison = comparison_table(q2,self.metrics).set_index("metric")
        for key in ("total_actual_cost_yuan","total_emergency_purchase_kwh","total_emergency_cost_yuan"):
            self.assertAlmostEqual(comparison.loc[key,"absolute_difference"],self.metrics[key]-q2[key])
            self.assertAlmostEqual(comparison.loc[key,"percentage_change"],(self.metrics[key]/q2[key]-1)*100)
        self.assertTrue(pd.isna(comparison.loc["simultaneous_slots","percentage_change"]))

    def test_paper_reports_from_saved_metrics(self):
        q2 = json.loads((ROOT/"outputs/q2/q2_metrics.json").read_text(encoding="utf-8"))
        with tempfile.TemporaryDirectory() as tmp:
            folder = Path(tmp)
            write_paper_reports(folder,dict(daily_metrics=self.daily),self.metrics,self.ranking,
                                self.selection,q2,"TEST ONLY",[])
            self.assertEqual(len(list(folder.iterdir())),6)
            updates = pd.read_csv(folder/"q3_update_time_summary.csv")
            self.assertEqual(updates.update_time.tolist(),["06:00","12:00","18:00"])
            self.assertAlmostEqual(updates.rho_mean.iloc[0],self.daily.rho_06.mean())
            table = pd.read_csv(folder/"q3_parameter_selection_table.csv")
            self.assertEqual(len(table),20)
            self.assertEqual(int(table.winner.sum()),1)
            text = (folder/"q3_paper_summary.md").read_text(encoding="utf-8")
            self.assertIn("TEST ONLY",text)
            self.assertIn("待确认事项：无",text)

    def test_template_mapping_format_and_raw_hash(self):
        template = ROOT/"data/raw/附件5/result3.xlsx"
        before = file_hash(template)
        inspect_template(template)
        dates = [str(d) for d in FORMAL_DATES]
        actual = pd.DataFrame(dict(date=np.repeat(dates,144),slot=np.tile(np.arange(1,145),334),
                                    x_real_kwh=1.,q_real_kwh=2.,z_real_kwh=3.,emergency_purchase_kwh=0.,
                                    soc_real_start_kwh=6000.,soc_real_end_kwh=6000.))
        actual.loc[[0,1,143,144],"emergency_purchase_kwh"] = 4.
        actual.loc[5,"emergency_purchase_kwh"] = 1e-7
        parts = []
        for hour in (0,6,12,18):
            slots = np.arange(hour*6+1,145)
            parts.append(pd.DataFrame(dict(date=np.repeat(dates,len(slots)),slot=np.tile(slots,334),
                                           update_time=f"{hour:02d}:00",new_plan_kwh=10.,increase_kwh=2. if hour else 0.,
                                           decrease_kwh=5. if hour else 0.,adjustment_cost_yuan=.5 if hour else 0.)))
        plans = pd.concat(parts,ignore_index=True)
        daily = pd.DataFrame(dict(date=dates,plan_cost_00_yuan=1440.))
        outputs = dict(actual_schedule=actual,plan_updates=plans,daily_metrics=daily)
        with tempfile.TemporaryDirectory() as tmp:
            dest = Path(tmp)/"synthetic_result3.xlsx"
            workbook = load_workbook(template)
            events = _populate(workbook,outputs)
            workbook.save(dest)
            workbook.close()
            self.assertEqual(len(events),3)
            verify_workbook(template,dest,outputs)
            written = load_workbook(dest,data_only=True)
            try:
                self.assertEqual(written.sheetnames,["计划购电量","调整购电量","充放电量","紧急购电量"])
                self.assertEqual(written["计划购电量"].cell(2,2).value,10.)
                self.assertEqual(written["调整购电量"].cell(2,2).value,0.)
                self.assertEqual(written["调整购电量"].cell(2,110).value,-9.)
                self.assertEqual(written["调整购电量"].cell(2,146).value,-648.)
                self.assertEqual(written["调整购电量"].cell(2,147).value,108.)
                self.assertEqual(written["充放电量"].max_row,2005)
                self.assertEqual(written["充放电量"].cell(2,3).value,72.)
                self.assertEqual(written["充放电量"].cell(2,4).value,72.)
                self.assertEqual(written["紧急购电量"].cell(2,3).value,8.)
                self.assertEqual(written["紧急购电量"].cell(3,3).value,4.)
                self.assertEqual(written["紧急购电量"].cell(4,3).value,4.)
            finally:
                written.close()
        self.assertEqual(before,file_hash(template))

    def test_submission_refuses_missing_final_results(self):
        with tempfile.TemporaryDirectory() as tmp:
            with self.assertRaises(FileNotFoundError):
                write_result3(Path(tmp))
            self.assertFalse((Path(tmp)/"outputs/submissions/result3.xlsx").exists())

    def test_readme_no_stale_search_status(self):
        text = (ROOT/"README.md").read_text(encoding="utf-8")
        self.assertNotIn("completed_groups=0",text)
        self.assertNotIn("尚无完整新参数组或最终排名",text)

    def test_saved_final_when_user_has_run_it(self):
        folder = ROOT/"outputs/q3/final"
        if not (folder/"q3_annual_metrics.json").exists():
            self.skipTest("最终334天由用户运行后再验收")
        outputs,metrics = load_final(folder)
        self.assertTrue(validate_final(outputs,metrics)["passed"])
        self.assertTrue(compare_search_winner(outputs["daily_metrics"],metrics,self.daily,self.metrics)["passed"])
        validation = json.loads((folder/"q3_final_validation.json").read_text(encoding="utf-8"))
        self.assertTrue(all(file_hash(ROOT/p)==h for p,h in validation["integrity"]["sha256"].items()))
        verify_workbook(ROOT/"data/raw/附件5/result3.xlsx",ROOT/"outputs/submissions/result3.xlsx",outputs)
