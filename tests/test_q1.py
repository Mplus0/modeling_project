"""Q1 独立约束复核、求解边界与模板写入测试。"""

from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from src.common.data_loader import discover_workbooks, file_hash
from src.q1.optimizer import VALIDATION_TOL, evaluate_schedule, load_input, solve_q1
from src.q1.result_writer import write_result

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "data/raw/附件5/result1.xlsx"


class Q1Tests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.frozen = discover_workbooks(ROOT / "data/raw") + sorted((ROOT / "data/processed").glob("*.csv"))
        cls.hashes = {p: file_hash(p) for p in cls.frozen}
        cls.frame = load_input(ROOT / "data/processed/q1_input.csv")
        cls.schedule, cls.metrics = solve_q1(cls.frame)

    def test_both_stages_optimal_and_objectives(self):
        s, m = self.schedule, self.metrics
        self.assertEqual(m["primary_status"], "optimal")
        self.assertEqual(m["secondary_status"], "optimal")
        cost = sum(s["price_yuan_per_kwh"] * (s["x_kwh"] + s["y_kwh"]))
        throughput = sum(0.9 * (s["x_kwh"] + s["q_kwh"]) + s["z_kwh"] / 0.9)
        self.assertAlmostEqual(cost, m["C_star"], delta=VALIDATION_TOL)
        self.assertAlmostEqual(cost, m["primary_cost"], delta=VALIDATION_TOL)
        self.assertAlmostEqual(throughput, m["secondary_solver_objective"], delta=VALIDATION_TOL)

    def test_all_constraints_independently(self):
        # 测试直接按模型公式计算，不以 evaluate_schedule 返回 true 代替验证。
        s = self.schedule
        self.assertEqual(len(s), 144)
        self.assertEqual(s["slot"].tolist(), list(range(1, 145)))
        self.assertFalse(s.isna().any().any())
        decisions = s[["x_kwh", "y_kwh", "q_kwh", "z_kwh"]].to_numpy()
        self.assertGreaterEqual(decisions.min(), -1e-6)
        soc = s["soc_kwh"].to_numpy()
        self.assertGreaterEqual(soc.min(), 1200 - 1e-6)
        self.assertLessEqual(soc.max(), 10800 + 1e-6)
        self.assertAlmostEqual(s["soc_previous_kwh"].iloc[0], 6000, delta=1e-6)
        self.assertAlmostEqual(soc[-1], 6000, delta=1e-6)
        np.testing.assert_allclose(soc, np.r_[6000, soc[:-1]] + 0.9 * (s["x_kwh"] + s["q_kwh"]) - s["z_kwh"] / 0.9,
                                   atol=1e-6, rtol=0)
        self.assertTrue((s["q_kwh"] <= s["pv_forecast_kwh"] + 1e-6).all())
        self.assertTrue((s["x_kwh"] + s["q_kwh"] <= 5000 / 6 + 1e-6).all())
        self.assertTrue((s["z_kwh"] <= 5000 / 6 + 1e-6).all())
        self.assertTrue((s["pv_forecast_kwh"] - s["q_kwh"] + s["y_kwh"] + s["z_kwh"] >= s["load_kwh"] - 1e-6).all())
        simultaneous = ((s["x_kwh"] + s["q_kwh"] > 1e-6) & (s["z_kwh"] > 1e-6)).sum()
        self.assertEqual(int(simultaneous), self.metrics["simultaneous_charge_discharge_intervals"])

    def test_secondary_removes_unnecessary_throughput(self):
        # 零负载、免费充裕光伏时，一级成本为零，二级应选择零吞吐；弃光允许。
        frame = self.frame.copy(deep=True)
        frame["price_yuan_per_kwh"] = 0.0
        frame["load_kw"] = 0.0
        frame["load_kwh"] = 0.0
        frame["pv_forecast_kw"] = 6000.0
        frame["pv_forecast_kwh"] = 1000.0
        schedule, metrics = solve_q1(frame)
        self.assertAlmostEqual(metrics["C_star"], 0, delta=1e-6)
        self.assertAlmostEqual(metrics["secondary_throughput"], 0, delta=1e-6)
        self.assertTrue(metrics["validation_passed"])
        self.assertTrue((schedule["supply_surplus_kwh"] >= 0).all())

    def test_invalid_schedule_is_detected(self):
        for column, value, residual in [("soc_kwh", 11000, "soc_bounds_kwh"),
                                         ("x_kwh", -1, "nonnegative_kwh"),
                                         ("z_kwh", 900, "discharge_limit_kwh")]:
            with self.subTest(column=column):
                corrupted = self.schedule.copy(deep=True)
                corrupted.loc[0, column] = value
                checked = evaluate_schedule(corrupted, self.metrics["C_star"], self.metrics["secondary_solver_objective"])
                self.assertFalse(checked["validation_passed"])
                self.assertGreater(checked["violations"][residual], 1e-6)

    def test_input_validation_and_immutability(self):
        original = self.frame.copy(deep=True)
        solve_q1(self.frame)
        pd.testing.assert_frame_equal(self.frame, original)
        for bad in (self.frame.iloc[:-1], self.frame.assign(load_kwh=np.nan), self.frame.assign(slot=1)):
            with self.assertRaises(ValueError):
                solve_q1(bad)
        self.assertEqual(self.hashes, {p: file_hash(p) for p in self.frozen})

    def test_default_mapping_follows_slot_row_order(self):
        with TemporaryDirectory() as directory:
            destination = Path(directory) / "result1.xlsx"
            write_result(self.schedule, self.metrics, TEMPLATE, destination)
            workbook = load_workbook(destination)
            try:
                for i in range(144):
                    self.assertAlmostEqual(workbook["计划购电量"].cell(i + 2, 2).value,
                                           self.schedule["grid_purchase_kwh"].iloc[i], delta=1e-6)
            finally:
                workbook.close()

    def test_writer_uses_explicit_mapping_and_preserves_template(self):
        before = file_hash(TEMPLATE)
        original = load_workbook(TEMPLATE)
        self.addCleanup(original.close)
        # 此处故意使用反序映射验证按标签写入；仅为临时测试夹具，不是正式时间决策。
        mapping = list(reversed([original["计划购电量"].cell(r, 1).value for r in range(2, 146)]))
        with TemporaryDirectory() as directory:
            destination = Path(directory) / "result1.xlsx"
            write_result(self.schedule, self.metrics, TEMPLATE, destination, mapping)
            result = load_workbook(destination)
            try:
                planned, battery = result["计划购电量"], result["充放电量"]
                values = {planned.cell(r, 1).value: planned.cell(r, 2).value for r in range(2, 146)}
                for i, label in enumerate(mapping):
                    self.assertAlmostEqual(values[label], self.schedule["grid_purchase_kwh"].iloc[i], delta=1e-6)
                for block in range(6):
                    part = self.schedule.iloc[block * 24:(block + 1) * 24]
                    self.assertAlmostEqual(battery.cell(block + 2, 2).value, (part["x_kwh"] + part["q_kwh"]).sum(), delta=1e-6)
                    self.assertAlmostEqual(battery.cell(block + 2, 3).value, part["z_kwh"].sum(), delta=1e-6)
                self.assertEqual(battery["E2"].value, 6000)
                self.assertEqual(battery["E3"].value, 6000)
                changed = {"计划购电量": {f"B{r}" for r in range(2, 146)},
                           "充放电量": {f"{c}{r}" for c in "BC" for r in range(2, 8)} | {"E2", "E3"}}
                for source in original:
                    target = result[source.title]
                    self.assertEqual(source.max_row, target.max_row)
                    self.assertEqual(source.max_column, target.max_column)
                    self.assertEqual(str(source.merged_cells), str(target.merged_cells))
                    self.assertEqual(source.page_setup, target.page_setup)
                    self.assertEqual(source.print_options, target.print_options)
                    self.assertEqual(source.sheet_format, target.sheet_format)
                    # 维度对象持有不同工作簿引用；比较实际序列化属性而非对象身份。
                    self.assertEqual({k: dict(v) for k, v in source.column_dimensions.items()},
                                     {k: dict(v) for k, v in target.column_dimensions.items()})
                    self.assertEqual({k: dict(v) for k, v in source.row_dimensions.items()},
                                     {k: dict(v) for k, v in target.row_dimensions.items()})
                    for row in source:
                        for cell in row:
                            # 未显式设样式的空白单元格与保存后的默认样式等价。
                            self.assertEqual(cell._style or original._cell_styles[0],
                                             target[cell.coordinate]._style or result._cell_styles[0])
                            if cell.coordinate not in changed[source.title]:
                                self.assertEqual(cell.value, target[cell.coordinate].value)
            finally:
                result.close()
        self.assertEqual(file_hash(TEMPLATE), before)

    def test_invalid_mapping_and_solution_rejected(self):
        with TemporaryDirectory() as directory:
            destination = Path(directory) / "result1.xlsx"
            with self.assertRaisesRegex(ValueError, "144"):
                write_result(self.schedule, self.metrics, TEMPLATE, destination, ["bad"] * 144)
            corrupted = self.schedule.copy(deep=True)
            corrupted.loc[143, "soc_kwh"] = 6001
            with self.assertRaisesRegex(ValueError, "约束"):
                write_result(corrupted, self.metrics, TEMPLATE, destination, [str(i) for i in range(144)])
            self.assertFalse(destination.exists())


if __name__ == "__main__":
    unittest.main()
