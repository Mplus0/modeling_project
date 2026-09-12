"""临时全年结构夹具测试写入和20组调度；不运行任何全年优化。"""
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from src.common.data_loader import file_hash
from src.q3.annual import FORMAL_DATES
from src.q3.parameters import PARAMETER_GRID
from src.q3.search import score_annual_results
from src.q4 import search
from src.q4.result_writer import render_copy,write_submission
from src.q4.q4_3 import REFERENCE_LABEL

ROOT = Path(__file__).resolve().parents[1]


def template_fixture(variant):
    dates = list(map(str,FORMAL_DATES))
    actual = pd.DataFrame(dict(date=np.repeat(dates,144),slot=np.tile(np.arange(1,145),334),
                              x_real_kwh=1.,q_real_kwh=2.,z_real_kwh=3.,emergency_purchase_kwh=0.,
                              soc_real_start_kwh=6000.,soc_real_end_kwh=6000.))
    for stage in ("primary","secondary","tertiary"):
        actual[f"rolling_{stage}_status"] = "optimal"
    actual.loc[[0,1,143,144],"emergency_purchase_kwh"] = 4.
    actual.loc[5,"emergency_purchase_kwh"] = 1e-7
    daily = pd.DataFrame(dict(date=dates,soc_start_kwh=6000.,soc_end_kwh=6000.,plan_cost_00_yuan=2880.,
                             actual_emergency_cost_yuan=0.,total_actual_cost_yuan=2880.,max_constraint_violation=0.))
    outputs = dict(actual_schedule=actual,daily_metrics=daily)
    if variant==2:
        outputs["plan_schedule"] = pd.DataFrame(dict(date=np.repeat(dates,144),slot=np.tile(np.arange(1,145),334),
                                                     grid_purchase_plan_kwh=10.,planned_cost_slot_yuan=10.,settled_plan_cost_yuan=20.))
    else:
        rows = []
        for hour in (0,6,12,18):
            slots = np.arange(hour*6+1,145)
            rows.append(pd.DataFrame(dict(date=np.repeat(dates,len(slots)),slot=np.tile(slots,334),update_time=f"{hour:02d}:00",
                                          new_plan_kwh=10.,increase_kwh=2. if hour else 0.,decrease_kwh=5. if hour else 0.,
                                          adjustment_cost_yuan=1. if hour else 0.)))
        outputs["plan_updates"] = pd.concat(rows,ignore_index=True)
    return outputs


class OutputTests(unittest.TestCase):
    def test_both_templates_copy_style_and_real_fee(self):
        for variant in (2,3):
            template = ROOT/f"data/raw/附件5/result4-{variant}.xlsx"
            before = file_hash(template)
            outputs = template_fixture(variant)
            with tempfile.TemporaryDirectory() as tmp:
                destination = Path(tmp)/"TEST_ONLY.xlsx"
                render_copy(template,destination,outputs,variant)
                raw,book = load_workbook(template),load_workbook(destination)
                try:
                    self.assertEqual(book.sheetnames,raw.sheetnames)
                    for name in raw.sheetnames:
                        self.assertEqual([c.value for c in raw[name][1]],[c.value for c in book[name][1]])
                        self.assertEqual(raw[name].column_dimensions["A"].width,book[name].column_dimensions["A"].width)
                        self.assertEqual(raw[name]["A1"]._style,book[name]["A1"]._style)
                    self.assertEqual(book["计划购电量"].cell(2,147).value,2880.)
                    self.assertEqual(book["充放电量"].max_row,2005)
                    self.assertEqual(raw["充放电量"].row_dimensions[2].height,book["充放电量"].row_dimensions[1994].height)
                    self.assertEqual(raw["充放电量"]["C2"]._style,book["充放电量"]["C1994"]._style)
                    self.assertEqual(book["紧急购电量"].max_row,4)
                    if variant==3:
                        self.assertEqual(book["调整购电量"].cell(2,146).value,-648.)
                        self.assertEqual(book["调整购电量"].cell(2,147).value,216.)
                finally:
                    raw.close()
                    book.close()
            self.assertEqual(file_hash(template),before)

    def test_submission_requires_human_review_and_final_data(self):
        with tempfile.TemporaryDirectory() as tmp:
            for variant in (2,3):
                with self.assertRaisesRegex(ValueError,"人工验收"):
                    write_submission(Path(tmp),variant)
                with self.assertRaises(FileNotFoundError):
                    write_submission(Path(tmp),variant,human_approved=True)
        self.assertIn("NOT FINAL",REFERENCE_LABEL)

    def test_cli_rejects_long_smoke_and_reference_overrides(self):
        for script,args in (("11_smoke_q4.py",["--days","334"]),("13_run_q4_3_reference.py",["--alpha","0.99"]),
                            ("13_run_q4_3_reference.py",["--lambda","1"])):
            result = subprocess.run([sys.executable,str(ROOT/"scripts"/script),*args],capture_output=True,text=True)
            self.assertEqual(result.returncode,2)
        self.assertFalse((ROOT/"outputs/submissions/result4-2.xlsx").exists())
        self.assertFalse((ROOT/"outputs/submissions/result4-3.xlsx").exists())

    def test_search_twenty_groups_resume_and_original_score(self):
        self.assertIs(search.score_annual_results,score_annual_results)
        self.assertEqual(len(PARAMETER_GRID),20)
        self.assertEqual({a for a,w in PARAMETER_GRID},{.8,.85,.9,.95})
        dates = list(map(str,FORMAL_DATES))
        calls = []
        def fake_run(history,prices,issues,variant,run_dates,label,alpha,weight,progress):
            calls.append((alpha,weight))
            self.assertEqual(tuple(run_dates),FORMAL_DATES)
            return dict(daily_metrics=pd.DataFrame(dict(date=dates,alpha=alpha,**{"lambda":weight},
                         total_actual_cost_yuan=100.+alpha+weight,emergency_slot_count=1,run_label=label))),0.
        def fake_write(folder,outputs,variant,runtime,label,integrity,formal=False):
            folder.mkdir(parents=True,exist_ok=True)
            frame = outputs["daily_metrics"]
            frame.to_csv(folder/"q4_daily_metrics.csv",index=False)
            metric = dict(alpha=float(frame.alpha.iloc[0]),**{"lambda":float(frame["lambda"].iloc[0])},
                          runtime_seconds=runtime,output_tables=["daily_metrics"])
            (folder/"q4_metrics.json").write_text(json.dumps(metric),encoding="utf-8")
            return metric
        with tempfile.TemporaryDirectory() as tmp, patch.object(search,"run_dates",side_effect=fake_run), \
                patch.object(search,"write_outputs",side_effect=fake_write),patch.object(search,"validate_outputs"):
            root = Path(tmp)
            ranking,winners = search.run_search(root,{},{},{},{})
            self.assertEqual(calls,list(PARAMETER_GRID))
            self.assertEqual(len(ranking),20)
            self.assertEqual(len(winners),1)
            search.run_search(root,{},{},{},{},resume=True)
            self.assertEqual(len(calls),20)
            group = root/"outputs/q4/search/groups/a0.80_l0.00/q4_daily_metrics.csv"
            group.write_text("corrupt",encoding="utf-8")
            with self.assertRaisesRegex(ValueError,"损坏"):
                search.run_search(root,{},{},{},{},resume=True)
