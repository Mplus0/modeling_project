"""仅复核已有年度CSV并写NOT_FOR_SUBMISSION临时副本，绝不调用正式导出入口。"""
from pathlib import Path
import json
import sys
import unittest
from datetime import datetime
from copy import copy

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
import numpy as np
from openpyxl import load_workbook
from src.common.data_loader import file_hash
from src.q4.metrics import load_outputs
from src.q4.validation import validate_outputs, VALIDATION_TOL
from src.q4.result_writer import (resolve_submission_source, validate_submission_source,
                                 render_copy, verify_values)
from src.q2.result_writer import slot_intervals, emergency_intervals

OUT = ROOT / "outputs/model_validation/final_export_precheck"


def save(name, value):
    (OUT/name).write_text(json.dumps(value,ensure_ascii=False,indent=2,allow_nan=False),encoding="utf-8")


def hashes():
    folders = ("data", "src/common", "src/q1", "src/q2", "src/q3", "src/q4",
               "outputs/q1", "outputs/q2", "outputs/q3/final", "outputs/q4", "outputs/submissions")
    return {p.relative_to(ROOT).as_posix():file_hash(p) for folder in folders
            for p in (ROOT/folder).rglob("*") if p.is_file() and "__pycache__" not in p.parts}


def template_check(template, trial, outputs, variant):
    def style(cell):
        # Excel序列化可能重排样式编号，比较实际属性而非内部索引。
        return (copy(cell.font),copy(cell.fill),copy(cell.border),copy(cell.alignment),
                copy(cell.protection),cell.number_format,cell.quotePrefix,cell.pivotButton)
    original,filled = load_workbook(template),load_workbook(trial)
    try:
        expected = ["计划购电量"] + (["调整购电量"] if variant==3 else []) + ["充放电量","紧急购电量"]
        assert original.sheetnames == filled.sheetnames == expected
        dimensions = {s.title:[s.max_row,s.max_column] for s in original}
        labels = [original[expected[0]].cell(1,c).value for c in range(2,146)]
        intervals = slot_intervals(labels)
        boundaries = {str(s):dict(column=s+1,template_label=labels[s-1],
                         model_start_minute=(s-1)*10,model_end_minute=s*10)
                      for s in (1,6,7,36,37,72,73,108,109,144)}
        # 保留此前确认的slot原行序；明确列出模板文字和模型时刻，不平移数值。
        for name in expected:
            a,b = original[name],filled[name]
            assert [c.value for c in a[1]] == [c.value for c in b[1]]
            assert a.max_column == b.max_column
            for x,y in zip(a[1],b[1]):
                assert style(x) == style(y), (name,x.coordinate,"header style")
            for col,dim in a.column_dimensions.items():
                assert dim.width == b.column_dimensions[col].width
        for name in expected[:-2]:
            a,b = original[name],filled[name]
            assert (a.max_row,a.max_column)==(335,147)==(b.max_row,b.max_column)
            assert str(a.merged_cells)==str(b.merged_cells)
            for r in range(2,336):
                assert a.cell(r,1).value == b.cell(r,1).value
                assert str(b.cell(r,1).value.date()) == outputs["daily_metrics"].date.iloc[r-2]
                for c in range(1,148):
                    assert style(a.cell(r,c)) == style(b.cell(r,c)), (name,r,c,"cell style")
        a,b = original["充放电量"],filled["充放电量"]
        assert b.max_row==2005
        for day in range(334):
            first = 2+day*6
            for offset in range(6):
                assert a.row_dimensions[2+offset].height == b.row_dimensions[first+offset].height
                for c in range(1,7):
                    x,y = a.cell(2+offset,c),b.cell(first+offset,c)
                    assert style(x)==style(y), ("充放电量",first+offset,c,"cell style")
                    if c in (2,5):
                        assert x.value==y.value
        expected_merges = set()
        from openpyxl.utils import get_column_letter
        for day in range(334):
            for m in a.merged_cells.ranges:
                if 2<=m.min_row<=m.max_row<=7:
                    expected_merges.add(f"{get_column_letter(m.min_col)}{m.min_row+day*6}:{get_column_letter(m.max_col)}{m.max_row+day*6}")
        assert set(map(str,b.merged_cells.ranges))==expected_merges
        events = emergency_intervals(outputs["actual_schedule"],intervals)
        a,b = original["紧急购电量"],filled["紧急购电量"]
        for i,event in enumerate(events):
            first = i==0 or events[i-1]["date"]!=event["date"]
            last = i==len(events)-1 or events[i+1]["date"]!=event["date"]
            source_row = 2 if first else (4 if last else 3)
            assert a.row_dimensions[source_row].height == b.row_dimensions[i+2].height
            for col in (1,2,3):
                x,y = a.cell(source_row,col),b.cell(i+2,col)
                if col==1:
                    # 现有writer将示例空白日期填为datetime，openpyxl自动赋予日期格式。
                    # 在内存复现该目标单元格行为；不修改原模板、不忽略其他格式属性。
                    from openpyxl.cell.cell import Cell
                    expected_date = Cell(a,row=source_row,column=col)
                    expected_date._style = copy(x._style)
                    expected_date.value = datetime.fromisoformat(event["date"])
                    x = expected_date
                assert style(x)==style(y), ("紧急购电量",source_row,i+2,col,
                    [(name,str(old),str(new)) for name,old,new in zip(
                        ("font","fill","border","alignment","protection","number_format","quotePrefix","pivotButton"),style(x),style(y)) if old!=new])
        errors = [(s.title,c.coordinate,c.value) for s in filled for row in s for c in row if c.data_type=="e"]
        assert not errors
        return dict(passed=True,sheet_names=expected,original_dimensions=dimensions,
                    output_dimensions={s.title:[s.max_row,s.max_column] for s in filled},
                    headers={s.title:[c.value for c in s[1]] for s in original},
                    original_formulas=[(s.title,c.coordinate,c.value) for s in original for row in s for c in row if c.data_type=="f"],
                    boundaries=boundaries,slot_mapping="沿用团队确认的slot原行序；模板标签保持不变",
                    update_first_slots={"00:00":1,"06:00":37,"12:00":73,"18:00":109},
                    units="购电、充放电、SOC均kWh；费用元；writer不再乘6或除6；价格元/kWh不单列",
                    existing_output_rule="variant2原浮点值；variant3复用_display，abs(value)<=1e-6显示0；紧急区间e>1e-6按日合并",
                    target_date_format="紧急表原General空白示例日期被填为datetime时，按既有openpyxl行为使用yyyy-mm-dd h:mm:ss；非目标格式保持",
                    mapping={"计划购电量!A2:A335":"date（原模板日期）",
                             "计划购电量!B:EO":"slot+1列；v2 grid_purchase_plan_kwh；v3 00:00 new_plan_kwh",
                             "计划购电量!EP":"144slot计划购电合计kWh",
                             "计划购电量!EQ":"plan_cost_00_yuan；v2 settled_plan_cost_yuan合计，真实价格结算",
                             "调整购电量(v3)!B:EO":"按slot汇总06/12/18 increase_kwh-decrease_kwh，非最终合同",
                             "调整购电量(v3)!EP/EQ":"净调整电量合计 / adjustment_cost_yuan合计",
                             "充放电量!C/D":"每天六行，每24slot合计x_real_kwh+q_real_kwh / z_real_kwh",
                             "充放电量!F":"每天首行初始SOC、次行最终SOC",
                             "紧急购电量!A/B/C":"date / 模板slot区间连续合并 / emergency_purchase_kwh合计",
                             "不额外写入":"动态单价、最终合同逐slot、实际总成本不增列，保留CSV"})
    finally:
        original.close()
        filled.close()


def main():
    OUT.mkdir(parents=True,exist_ok=True)
    # 每次预检重新执行门槛测试，不能把上次测试记录当成当前通过证据。
    save("final_export_precheck.json",dict(status="BLOCKED",reason="precheck进行中，尚未完成全部验收"))
    suite = unittest.defaultTestLoader.discover(str(ROOT/"tests"),pattern="test_q4_submission_gate.py")
    tested = unittest.TextTestRunner(verbosity=1).run(suite)
    tests = dict(tests=tested.testsRun,failures=len(tested.failures),errors=len(tested.errors),passed=tested.wasSuccessful())
    save("tests.json",tests)
    if not tested.wasSuccessful():
        return
    before = hashes()
    save("adaptation_frozen_before.json",before)
    reports,templates = {},{}
    for variant in (2,3):
        print(f"Q4-{variant} read-only validation",flush=True)
        report = dict(variant=variant,errors=[])
        folder = resolve_submission_source(ROOT,variant)
        outputs,metrics = load_outputs(folder)
        report["source"] = folder.relative_to(ROOT).as_posix()
        checked = validate_outputs(outputs,variant,formal=True)
        report.update(checked)
        report["metrics"] = {k:v for k,v in metrics.items() if k!="protected_sha256"}
        daily = outputs["daily_metrics"]
        recomputed = dict(total_actual_cost_yuan=float(daily.total_actual_cost_yuan.sum()),
                          total_plan_cost_real_yuan=float(daily.plan_cost_00_yuan.sum()),
                          total_adjustment_cost_real_yuan=float(daily.adjustment_cost_total_yuan.sum()),
                          total_emergency_purchase_kwh=float(daily.actual_emergency_purchase_kwh.sum()),
                          total_emergency_cost_yuan=float(daily.actual_emergency_cost_yuan.sum()),
                          emergency_slot_count=int(daily.emergency_slot_count.sum()),
                          emergency_day_count=int((daily.emergency_slot_count>0).sum()),
                          initial_soc_kwh=float(daily.soc_start_kwh.iloc[0]),final_soc_kwh=float(daily.soc_end_kwh.iloc[-1]),
                          minimum_soc_kwh=float(daily.soc_min_kwh.min()),maximum_soc_kwh=float(daily.soc_max_kwh.max()),
                          simultaneous_slots=int(daily.simultaneous_slots.sum()))
        report["metric_recalculation_differences"]={k:abs(v-metrics[k]) for k,v in recomputed.items()}
        assert max(report["metric_recalculation_differences"].values())<=VALIDATION_TOL
        report["nan_count"] = sum(int(f.isna().sum().sum()) for f in outputs.values())
        report["inf_count"] = sum(int(np.isinf(f.select_dtypes(include=[np.number])).sum().sum()) for f in outputs.values())
        actual = outputs["actual_schedule"]
        nonnegative = [c for c in actual if c.startswith(("x_real","y_real","q_real","z_real","soc_real","grid_purchase","emergency_purchase"))]
        report["nonnegative_minima"] = {c:float(actual[c].min()) for c in nonnegative}
        assert min(report["nonnegative_minima"].values()) >= -VALIDATION_TOL
        assert report["nan_count"] == report["inf_count"] == 0
        plan = outputs["plan_schedule" if variant==2 else "plan_updates"]
        columns = (["grid_purchase_plan_kwh","x_plan_kwh","y_plan_kwh","q_plan_kwh","z_plan_kwh"] if variant==2
                   else ["new_plan_kwh","previous_plan_kwh","increase_kwh","decrease_kwh"])
        report["plan_nonnegative_minima"]={c:float(plan[c].min()) for c in columns if c in plan}
        assert min(report["plan_nonnegative_minima"].values())>=-VALIDATION_TOL
        try:
            validate_submission_source(ROOT,variant,outputs,metrics)
            report["source_gate_passed"] = True
        except Exception as error:
            report["source_gate_passed"] = False
            report["errors"].append(str(error))
        mismatches = {p:dict(expected=h,actual=file_hash(ROOT/p) if (ROOT/p).is_file() else None)
                      for p,h in metrics["protected_sha256"].items()
                      if not (ROOT/p).is_file() or file_hash(ROOT/p)!=h}
        report["protected_sha_mismatches"] = mismatches
        # 临时副本仅用于完整数值/样式验证，不调用write_submission。
        trial_dir = OUT/"tmp"/datetime.now().strftime("%Y%m%d_%H%M%S_%f")
        trial = trial_dir/f"NOT_FOR_SUBMISSION_result4-{variant}.xlsx"
        template = ROOT/f"data/raw/附件5/result4-{variant}.xlsx"
        try:
            rendered = render_copy(template,trial,outputs,variant)
            verify_values(trial,outputs,variant)
            report["temporary_excel"] = str(trial.relative_to(ROOT))
            report["verify_values_passed"] = True
            templates[str(variant)] = template_check(template,trial,outputs,variant)
            templates[str(variant)].update(rendered)
        except Exception as error:
            report["errors"].append(f"模板模拟验收：{type(error).__name__}: {error}")
            templates[str(variant)] = dict(passed=False,error=repr(error))
        report["status"] = "BLOCKED" if report["errors"] else "READY_FOR_HUMAN_APPROVAL"
        reports[str(variant)] = report
        save(f"q4_{variant}_precheck.json",report)
    after = hashes()
    save("adaptation_frozen_after.json",after)
    integrity = dict(unchanged=before==after,changed=[p for p in before.keys()|after.keys() if before.get(p)!=after.get(p)],checked_files=len(after))
    save("adaptation_integrity.json",integrity)
    save("template_precheck.json",templates)
    tests = json.loads((OUT/"tests.json").read_text(encoding="utf-8"))
    result = dict(status="READY_FOR_HUMAN_APPROVAL" if integrity["unchanged"] and all(r["status"]=="READY_FOR_HUMAN_APPROVAL" for r in reports.values()) else "BLOCKED",
                  variants={k:r["status"] for k,r in reports.items()},integrity=integrity,
                  formal_export_executed=False,solver_calls=0,
                  existing_result4=[v for v in (2,3) if (ROOT/f"outputs/submissions/result4-{v}.xlsx").exists()])
    if result["existing_result4"]:
        result["status"]="BLOCKED"
    result["targeted_tests"]=tests
    if not tests.get("passed"):
        result["status"]="BLOCKED"
    save("writer_precheck.json",dict(source_2="outputs/q4/q4_2",source_3="outputs/q4/reference/q4_3_a0.85_l0.25",
         gate="team-confirmed reference，formal及SMOKE校验、年度validate_outputs、逐文件protected_sha256",
         calls_optimizer=False,requires_search_summary=False,rejects_existing_destination=True,
         audits="outputs/model_validation/final_export/q4_{variant}_submission_validation.json",
         raw_template_saved=False,formal_export_executed=False))
    save("final_export_precheck.json",result)
    text = ["# Q4最终导出前只读验收", "", f"状态：{result['status']}",
            "", "未调用优化器，未生成正式result4。临时文件为NOT_FOR_SUBMISSION，仅用于验收。"]
    for variant,r in reports.items():
        text += ["",f"## Q4-{variant}",f"来源：{r['source']}",f"年度数值验收：{r['passed']}；{r['days']}天、{r['slots']}slot；最大误差{r['max_constraint_violation']}；跨日SOC误差{r['cross_day_soc_max_difference']}",
                 f"NaN={r['nan_count']}，Inf={r['inf_count']}；临时Excel verify_values={r.get('verify_values_passed',False)}。",f"来源门槛：{r['source_gate_passed']}"]
        text += [f"- {error}" for error in r["errors"]]
    text += ["",f"本轮冻结SHA一致：{integrity['unchanged']}。逐项结果见JSON。历史protected SHA未擅自迁移或忽略。"]
    text += ["", "时间映射沿用已确认的slot原行序，不宣称模板标签与模型时钟相同：模型slot1为00:00–00:10，官方首列仍为0:10–0:20；06/12/18更新从slot37/73/109生效。边界对照见template_precheck.json，未移动数值或改写官方标签。",
             "紧急表原空白日期单元格General在写入datetime时按既有writer行为获得日期格式；其他字体、边框等属性已核验。未修改render_copy或verify_values。"]
    (OUT/"final_export_precheck.md").write_text("\n".join(text),encoding="utf-8")
    print(json.dumps(result,ensure_ascii=False),flush=True)


if __name__=="__main__":
    main()
