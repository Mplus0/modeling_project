"""复制并填写官方 Q1 模板；默认按用户确认的 slot 行序写入。"""

from pathlib import Path
from shutil import copy2
from tempfile import TemporaryDirectory

from openpyxl import load_workbook

from src.common.data_loader import file_hash
from src.q1.optimizer import evaluate_schedule


def write_result(schedule, metrics, template, destination, interval_mapping=None):
    """默认 slot i 写入模板第 i+1 行；可显式传入已确认的区间标签列表。"""
    checked = evaluate_schedule(schedule, metrics["C_star"], metrics["secondary_solver_objective"])
    if not checked["validation_passed"] or any(metrics[k] != "optimal" for k in ("primary_status", "secondary_status")):
        raise ValueError("调度未通过两阶段最优性及独立约束检查，拒绝写入模板")
    template, destination = Path(template), Path(destination)
    if template.resolve() == destination.resolve() or destination.is_symlink():
        raise ValueError("结果路径不得覆盖或链接到原始模板")
    if interval_mapping is not None and (len(interval_mapping) != 144 or len(set(interval_mapping)) != 144):
        raise ValueError("正式时间映射必须包含 144 个唯一模板区间")
    before = file_hash(template)
    destination.parent.mkdir(parents=True, exist_ok=True)
    with TemporaryDirectory(prefix=".q1-result-", dir=destination.parent) as temporary:
        staged = Path(temporary) / "result1.xlsx"
        # 必须复制原模板后仅修改目标单元格；不重建工作簿或改动格式、标签。
        copy2(template, staged)
        workbook = load_workbook(staged)
        try:
            if workbook.sheetnames != ["计划购电量", "充放电量"]:
                raise ValueError("Q1 官方模板工作表不符")
            planned, battery = workbook["计划购电量"], workbook["充放电量"]
            if (planned.max_row, planned.max_column) != (145, 2) or (battery.max_row, battery.max_column) != (7, 5):
                raise ValueError("Q1 官方模板尺寸不符")
            rows = {planned.cell(r, 1).value: r for r in range(2, 146)}
            # 用户已确认按 slot 行序填写数值；保留官方区间标签，不作十分钟平移。
            if interval_mapping is None:
                interval_mapping = [planned.cell(r, 1).value for r in range(2, 146)]
            if len(rows) != 144:
                raise ValueError("官方模板购电区间标签不唯一")
            if set(rows) != set(interval_mapping):
                raise ValueError("时间映射与官方模板区间不一致")
            for slot, label in enumerate(interval_mapping):
                planned.cell(rows[label], 2, float(schedule["x_kwh"].iloc[slot] + schedule["y_kwh"].iloc[slot]))
            expected_blocks = [f"{h}:00-{h+4}:00" for h in range(0, 24, 4)]
            if [battery.cell(r, 1).value for r in range(2, 8)] != expected_blocks:
                raise ValueError("Q1 四小时模板标签不符")
            if [battery.cell(r, 4).value for r in (2, 3)] != ["0:00", "24:00"]:
                raise ValueError("Q1 初末储电量标签不符")
            # 用户已确认固定连续 24 个 slot 汇总；此处用外部输入/送达端电量。
            for block in range(6):
                part = schedule.iloc[block * 24:(block + 1) * 24]
                battery.cell(block + 2, 2, float((part["x_kwh"] + part["q_kwh"]).sum()))
                battery.cell(block + 2, 3, float(part["z_kwh"].sum()))
            battery.cell(2, 5, float(schedule["soc_previous_kwh"].iloc[0]))
            battery.cell(3, 5, float(schedule["soc_kwh"].iloc[-1]))
            workbook.save(staged)
        finally:
            workbook.close()
        if file_hash(template) != before:
            raise RuntimeError("官方模板哈希发生变化")
        staged.replace(destination)
    return destination
