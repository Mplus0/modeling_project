"""按官方示例样式扩展输出副本，绝不修改原始result2模板。"""

from copy import copy
from datetime import datetime
from pathlib import Path
from shutil import copy2
from tempfile import TemporaryDirectory

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from src.common.data_loader import file_hash
from src.common.time_utils import clock_seconds
from src.q2.forecasting import FORMAL_DATES
from src.q2.optimizer import VALIDATION_TOL

# 官方题面表3指定日期、表1指定起始时刻；所有数值均从CSV自动提取。
PAPER_DATES = ("2025-03-20", "2025-06-21", "2025-09-23", "2025-12-21")
PAPER_HOURS = (10, 12, 14, 16, 18, 20)


def slot_intervals(labels):
    intervals = []
    for i, label in enumerate(labels):
        parts = str(label).split("-")
        if len(parts) != 2:
            raise ValueError(f"无法解析官方购电区间 {label}")
        start, end = [clock_seconds(p) for p in parts]
        if start is None or end is None:
            raise ValueError(f"官方购电区间时刻无效 {label}")
        # 沿用Q1的slot行序，午夜标记按模板连续顺序展开，不重新平移10分钟。
        if intervals and start < intervals[-1][0]:
            start += 86400
        if end < start:
            end += 86400
        if end - start != 600 or (intervals and start != intervals[-1][1]):
            raise ValueError("官方购电区间不连续或不为10分钟")
        intervals.append((start, end))
    if len(intervals) != 144:
        raise ValueError("官方模板必须有144个购电区间")
    return intervals


def format_clock(seconds):
    day, within = divmod(int(seconds), 86400)
    hour, minute = divmod(within // 60, 60)
    return f"{hour}:{minute:02d}" + (f"+{day}" if day else "")


def emergency_intervals(actual, intervals, tolerance=VALIDATION_TOL):
    records = []
    for day, frame in actual.groupby("date", sort=True):
        if frame.slot.tolist() != list(range(1, 145)):
            raise ValueError(f"{day} 实际slot不完整或乱序")
        values = frame.emergency_purchase_kwh.to_numpy(float)
        if not np.isfinite(values).all() or values.min() < -tolerance:
            raise ValueError("紧急购电包含非法数值")
        selected = np.flatnonzero(values > tolerance)
        if not len(selected):
            continue
        groups = np.split(selected, np.flatnonzero(np.diff(selected) != 1) + 1)
        # 每天独立分组，次日第一个slot不能与前一天的最后一个slot合并。
        for group in groups:
            start, end = int(group[0]), int(group[-1])
            records.append({"date": str(day), "start_slot": start + 1, "end_slot": end + 1,
                            "interval": f"{format_clock(intervals[start][0])}-{format_clock(intervals[end][1])}",
                            "emergency_purchase_kwh": float(values[group].sum())})
    return records


def capture_block(sheet, first, last):
    cells = [[(cell.value, copy(cell._style), copy(cell.comment), copy(cell.hyperlink)) for cell in row]
             for row in sheet.iter_rows(min_row=first, max_row=last)]
    dimensions = [copy(sheet.row_dimensions[r]) for r in range(first, last + 1)]
    merges = [(r.min_row - first, r.max_row - first, r.min_col, r.max_col)
              for r in sheet.merged_cells.ranges if first <= r.min_row <= r.max_row <= last]
    return cells, dimensions, merges


def paste_block(sheet, block, start):
    cells, dimensions, merges = block
    for offset, row in enumerate(cells):
        dimension = copy(dimensions[offset])
        dimension.index = start + offset
        sheet.row_dimensions[start + offset] = dimension
        for col, (value, style, comment, hyperlink) in enumerate(row, 1):
            cell = sheet.cell(start + offset, col, value)
            cell._style, cell.comment, cell.hyperlink = copy(style), copy(comment), copy(hyperlink)
    for first, last, left, right in merges:
        sheet.merge_cells(start_row=start + first, end_row=start + last, start_column=left, end_column=right)


def clear_body(sheet):
    for merged in list(sheet.merged_cells.ranges):
        if merged.min_row >= 2:
            sheet.unmerge_cells(str(merged))
    sheet.delete_rows(2, sheet.max_row - 1)
    for row in list(sheet.row_dimensions):
        if row >= 2:
            del sheet.row_dimensions[row]


def validate_tables(plan, actual, daily):
    expected = [d.isoformat() for d in FORMAL_DATES]
    for name, frame in [("计划", plan), ("实际", actual)]:
        if len(frame) != 48096 or frame.date.drop_duplicates().tolist() != expected:
            raise ValueError(f"Q2{name}输出日期或行数不完整")
        if not all(g.slot.tolist() == list(range(1, 145)) for _, g in frame.groupby("date", sort=False)):
            raise ValueError(f"Q2{name}slot顺序不符")
    if daily.date.tolist() != expected:
        raise ValueError("每日指标日期不完整")
    status = ["plan_primary_status", "plan_secondary_status", "actual_primary_status", "actual_secondary_status"]
    if not daily[status].eq("optimal").all().all():
        raise ValueError("Q2存在非最优日期，拒绝生成提交文件")
    if daily[["max_plan_constraint_violation", "max_actual_constraint_violation"]].max().max() > VALIDATION_TOL:
        raise ValueError("Q2存在约束验证失败日期")


def paper_tables(plan, actual, daily, labels, intervals, emergencies):
    lines = ["# 问题二论文表格", "", "单位：电量 kWh，费用元。采用已确认的slot顺序；购电时段沿用官方模板标签。",
             "充放电按连续24个slot汇总，使用实际 x_real+q_real 和 z_real。全天实际费用包含计划费用与5倍分时紧急费用。", ""]
    for day in PAPER_DATES:
        planned, operated = plan[plan.date == day], actual[actual.date == day]
        metric = daily[daily.date == day].iloc[0]
        lines += [f"## {day}", "", "| 指定购电时段 | 计划购电量 |", "|---|---:|"]
        for hour in PAPER_HOURS:
            slot = next(i for i, (start, _) in enumerate(intervals) if start == hour * 3600)
            lines.append(f"| {labels[slot]} | {planned.grid_purchase_plan_kwh.iloc[slot]:.8f} |")
        lines += ["", f"全天计划购电量：{metric.planned_grid_purchase_kwh:.8f}；计划费用：{metric.planned_cost_yuan:.8f}；紧急费用：{metric.emergency_cost_yuan:.8f}；全天实际总费用：{metric.total_cost_yuan:.8f}。", "",
                  "| 四小时时段 | 实际充电量 | 实际放电量 |", "|---|---:|---:|"]
        for block in range(6):
            part = operated.iloc[block * 24:(block + 1) * 24]
            lines.append(f"| {block*4}:00-{(block+1)*4}:00 | {(part.x_real_kwh+part.q_real_kwh).sum():.8f} | {part.z_real_kwh.sum():.8f} |")
        lines += ["", f"0:00实际SOC：{metric.soc_start_kwh:.8f}；24:00实际SOC：{metric.soc_actual_end_kwh:.8f}。", "",
                  "| 紧急购电区间 | 紧急购电量 |", "|---|---:|"]
        events = [r for r in emergencies if r["date"] == day]
        lines.extend(f"| {r['interval']} | {r['emergency_purchase_kwh']:.8f} |" for r in events)
        if not events:
            lines.append("| 无超过报告容差的紧急购电 | 0 |")
        lines.append("")
    return "\n".join(lines)


def write_outputs(template, destination, plan, actual, daily, paper_path):
    validate_tables(plan, actual, daily)
    template, destination, paper_path = map(Path, (template, destination, paper_path))
    if destination.resolve() == template.resolve() or destination.is_symlink():
        raise ValueError("输出副本不得覆盖原始模板")
    before = file_hash(template)
    destination.parent.mkdir(parents=True, exist_ok=True)
    with TemporaryDirectory(prefix=".q2-result-", dir=destination.parent) as temporary:
        staged = Path(temporary) / "result2.xlsx"
        copy2(template, staged)
        workbook = load_workbook(staged)
        try:
            if workbook.sheetnames != ["计划购电量", "充放电量", "紧急购电量"]:
                raise ValueError("Q2模板工作表不符")
            planned_sheet, battery, emergency = [workbook[s] for s in workbook.sheetnames]
            if (planned_sheet.max_row, planned_sheet.max_column) != (335, 147):
                raise ValueError("计划购电模板尺寸不符")
            labels = [planned_sheet.cell(1, col).value for col in range(2, 146)]
            intervals = slot_intervals(labels)
            events = emergency_intervals(actual, intervals)
            charge_block = capture_block(battery, 2, 7)
            emergency_styles = [capture_block(emergency, r, r) for r in (2, 3, 4)]
            clear_body(battery)
            clear_body(emergency)
            for index, day in enumerate(FORMAL_DATES):
                day_text = day.isoformat()
                p = plan.iloc[index * 144:(index + 1) * 144]
                a = actual.iloc[index * 144:(index + 1) * 144]
                if pd.Timestamp(planned_sheet.cell(index + 2, 1).value).date() != day:
                    raise ValueError("计划购电模板日期不一致")
                for slot, value in enumerate(p.grid_purchase_plan_kwh, 2):
                    planned_sheet.cell(index + 2, slot, float(value))
                planned_sheet.cell(index + 2, 146, float(p.grid_purchase_plan_kwh.sum()))
                # 计划工作表的费用为计划购电费用；实际总费用另见每日指标和论文表。
                planned_sheet.cell(index + 2, 147, float(p.planned_cost_slot_yuan.sum()))
                first = 2 + index * 6
                paste_block(battery, charge_block, first)
                battery.cell(first, 1, datetime.combine(day, datetime.min.time()))
                for block in range(6):
                    part = a.iloc[block * 24:(block + 1) * 24]
                    battery.cell(first + block, 3, float((part.x_real_kwh + part.q_real_kwh).sum()))
                    battery.cell(first + block, 4, float(part.z_real_kwh.sum()))
                battery.cell(first, 6, float(a.soc_real_start_kwh.iloc[0]))
                battery.cell(first + 1, 6, float(a.soc_real_end_kwh.iloc[-1]))
            for i, event in enumerate(events):
                first_of_day = i == 0 or events[i - 1]["date"] != event["date"]
                last_of_day = i == len(events) - 1 or events[i + 1]["date"] != event["date"]
                style = emergency_styles[0 if first_of_day else (2 if last_of_day else 1)]
                paste_block(emergency, style, i + 2)
                # 每个区间显式写日期以便独立复算，不填没有紧急购电的日期。
                emergency.cell(i + 2, 1, datetime.fromisoformat(event["date"]))
                emergency.cell(i + 2, 2, event["interval"])
                emergency.cell(i + 2, 3, event["emergency_purchase_kwh"])
            workbook.save(staged)
            text = paper_tables(plan, actual, daily, labels, intervals, events)
        finally:
            workbook.close()
        if file_hash(template) != before:
            raise RuntimeError("原始Q2模板SHA-256发生变化")
        staged.replace(destination)
    paper_path.parent.mkdir(parents=True, exist_ok=True)
    paper_path.write_text(text, encoding="utf-8")
    return events
