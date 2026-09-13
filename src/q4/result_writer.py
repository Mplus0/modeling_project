"""仅在人工验收后填充官方Q4副本；复用Q2/Q3样式和已确认字段映射。"""

import json
from pathlib import Path
from shutil import copy2
from tempfile import TemporaryDirectory
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from src.common.data_loader import file_hash
from src.q2.result_writer import write_outputs as write_two, emergency_intervals, slot_intervals
from src.q3.result_writer import _populate, verify_workbook
from src.q4.metrics import load_outputs
from src.q4.validation import validate_outputs
from src.q2.optimizer import VALIDATION_TOL

REFERENCE_SOURCE = "outputs/q4/reference/q4_3_a0.85_l0.25"
ADOPTION = "团队确认第四问固定沿用第三问确定的风险参数，不再进行Q4参数重新搜索。"


def resolve_submission_source(root, variant):
    root = Path(root)
    if variant == 2:
        return root / "outputs/q4/q4_2"
    if variant != 3:
        raise ValueError("Q4模板类型必须为2或3")
    acceptance = json.loads((root / "outputs/model_validation/q4_reference_acceptance.json").read_text(encoding="utf-8"))
    required = dict(passed=True, team_confirmed=True, alpha=.85, source=REFERENCE_SOURCE,
                    label="Q4-3 fixed-parameter final candidate", parameter_adoption=ADOPTION)
    required["lambda"] = .25
    for key, value in required.items():
        if acceptance.get(key) != value or (isinstance(value, bool) and acceptance.get(key) is not value):
            raise ValueError(f"Q4团队确认记录不符：{key}")
    return root / REFERENCE_SOURCE


def check_protected(root, metrics):
    # 保留原逐文件SHA门槛；不迁移或忽略历史冻结记录。
    hashes = metrics.get("protected_sha256")
    if not isinstance(hashes, dict) or not hashes:
        raise ValueError("Q4缺少protected_sha256")
    changed = [p for p,h in hashes.items() if not (root/p).is_file() or file_hash(root/p) != h]
    if changed:
        raise ValueError("Q4冻结文件变化：" + ", ".join(changed))


def validate_submission_source(root, variant, outputs, metrics):
    resolve_submission_source(root, variant)
    if metrics.get("formal") is not True or metrics.get("variant") != variant:
        raise ValueError("Q4 formal/variant不符")
    label = str(metrics.get("run_label", "")).upper()
    if "SMOKE" in label or (variant == 2 and "REFERENCE" in label):
        raise ValueError("Q4运行标签不允许正式导出")
    if variant == 3 and (metrics.get("alpha") != .85 or metrics.get("lambda") != .25):
        raise ValueError("Q4 reference参数不符")
    checked = validate_outputs(outputs, variant, formal=True)
    if variant == 3:
        acceptance = json.loads((root/"outputs/model_validation/q4_reference_acceptance.json").read_text(encoding="utf-8"))
        # 采用资格来自团队记录，历史REFERENCE标签原样保留；年度数值仍须逐项核验。
        for record in (metrics, acceptance, checked):
            if record.get("passed") is not True or record.get("days") != 334 or record.get("slots") != 48096:
                raise ValueError("Q4年度验收范围不符")
            for key in ("max_constraint_violation", "cross_day_soc_max_difference"):
                value = record.get(key, float("nan"))
                if not np.isfinite(value) or not 0 <= value <= VALIDATION_TOL:
                    raise ValueError(f"Q4年度验收超差：{key}")
        aliases = {"plan_cost_yuan":"total_plan_cost_real_yuan", "adjustment_cost_yuan":"total_adjustment_cost_real_yuan",
                   "emergency_purchase_kwh":"total_emergency_purchase_kwh", "emergency_cost_yuan":"total_emergency_cost_yuan"}
        for key in (*aliases,"total_actual_cost_yuan","emergency_slot_count","emergency_day_count",
                    "initial_soc_kwh","final_soc_kwh","minimum_soc_kwh","maximum_soc_kwh","simultaneous_slots"):
            if key not in acceptance or aliases.get(key,key) not in metrics:
                raise ValueError(f"Q4验收缺少交叉核对指标：{key}")
        for key, value in acceptance.items():
            target = aliases.get(key, key)
            if isinstance(value, (int,float)) and not isinstance(value,bool) and target in metrics:
                if not np.isfinite(metrics[target]) or abs(value-metrics[target]) > VALIDATION_TOL:
                    raise ValueError(f"Q4 acceptance与metrics冲突：{key}")
    check_protected(root, metrics)
    return checked


def render_copy(template, destination, outputs, variant):
    """底层模板复制器供临时合成测试；正式调用必须经过write_submission门槛。"""
    template,destination = Path(template),Path(destination)
    if template.resolve()==destination.resolve() or destination.is_symlink():
        raise ValueError("Q4输出不得覆盖模板")
    before = file_hash(template)
    actual,daily = outputs["actual_schedule"],outputs["daily_metrics"]
    destination.parent.mkdir(parents=True,exist_ok=True)
    with TemporaryDirectory(dir=destination.parent,prefix=".q4-template-") as tmp:
        staged = Path(tmp)/destination.name
        if variant==3:
            copy2(template,staged)
            book = load_workbook(staged)
            try:
                _populate(book,outputs)
                book.save(staged)
            finally:
                book.close()
            verify_workbook(template,staged,outputs)
        elif variant==2:
            # 仅写入视图采用真实结算费用，不修改计划CSV中的预测价格或目标。
            plan = outputs["plan_schedule"].copy()
            plan["planned_cost_slot_yuan"] = plan.settled_plan_cost_yuan
            view = daily.rename(columns={"soc_end_kwh":"soc_actual_end_kwh","plan_cost_00_yuan":"planned_cost_yuan",
                                         "actual_emergency_cost_yuan":"emergency_cost_yuan","total_actual_cost_yuan":"total_cost_yuan"}).copy()
            for stage in ("primary","secondary"):
                view[f"plan_{stage}_status"] = "optimal"
            for name in ("max_plan_constraint_violation","max_actual_constraint_violation","max_rolling_constraint_violation"):
                view[name] = daily.max_constraint_violation.to_numpy()
            view["planned_grid_purchase_kwh"] = plan.groupby("date",sort=False).grid_purchase_plan_kwh.sum().to_numpy()
            write_two(template,staged,plan,actual,view,Path(tmp)/"unused_paper.md")
        else:
            raise ValueError("Q4模板类型必须为2或3")
        verify_values(staged,outputs,variant)
        if file_hash(template)!=before:
            raise RuntimeError("Q4原模板哈希变化")
        staged.replace(destination)
    return dict(template_sha256=before,output_sha256=file_hash(destination),passed=True)


def verify_values(path, outputs, variant):
    """由CSV重算每日费用、24slot充放电和紧急区间，不依赖写入函数的返回值。"""
    book = load_workbook(path,data_only=True)
    try:
        names = ["计划购电量","充放电量","紧急购电量"] if variant==2 else ["计划购电量","调整购电量","充放电量","紧急购电量"]
        if book.sheetnames!=names or book["充放电量"].max_row!=2005:
            raise ValueError("Q4模板sheet或全年充放电结构不符")
        actual,daily = outputs["actual_schedule"],outputs["daily_metrics"]
        for i,row in enumerate(daily.itertuples()):
            a = actual[actual.date.eq(row.date)]
            grid = (outputs["plan_schedule"].loc[lambda f:f.date.eq(row.date),"grid_purchase_plan_kwh"].to_numpy()
                    if variant==2 else outputs["plan_updates"].loc[lambda f:f.date.eq(row.date)&f.update_time.eq("00:00"),"new_plan_kwh"].to_numpy())
            values = [book["计划购电量"].cell(i+2,col).value for col in range(2,148)]
            np.testing.assert_allclose(values,np.r_[grid,grid.sum(),row.plan_cost_00_yuan],atol=1e-6,rtol=0)
            for block in range(6):
                part = a.iloc[block*24:(block+1)*24]
                np.testing.assert_allclose([book["充放电量"].cell(2+i*6+block,c).value for c in (3,4)],
                                            [(part.x_real_kwh+part.q_real_kwh).sum(),part.z_real_kwh.sum()],atol=1e-6,rtol=0)
            np.testing.assert_allclose([book["充放电量"].cell(2+i*6,6).value,book["充放电量"].cell(3+i*6,6).value],
                                        [row.soc_start_kwh,row.soc_end_kwh],atol=1e-6,rtol=0)
            if variant==3:
                changes = outputs["plan_updates"].loc[lambda f:f.date.eq(row.date)&~f.update_time.eq("00:00")]
                net = (changes.increase_kwh-changes.decrease_kwh).groupby(changes.slot).sum().reindex(range(1,145),fill_value=0.)
                np.testing.assert_allclose([book["调整购电量"].cell(i+2,c).value for c in range(2,148)],
                                            np.r_[net,net.sum(),changes.adjustment_cost_yuan.sum()],atol=1e-6,rtol=0)
        labels = [book["计划购电量"].cell(1,c).value for c in range(2,146)]
        events = emergency_intervals(actual,slot_intervals(labels))
        if book["紧急购电量"].max_row!=len(events)+1:
            raise ValueError("Q4紧急购电区间数不符")
        for i,event in enumerate(events,2):
            cells = [book["紧急购电量"].cell(i,c).value for c in (1,2,3)]
            if str(cells[0].date())!=event["date"] or cells[1]!=event["interval"] or abs(cells[2]-event["emergency_purchase_kwh"])>1e-6:
                raise ValueError("Q4紧急区间无法由实际e独立复算")
    finally:
        book.close()


def write_submission(root, variant, human_approved=False):
    if not human_approved:
        raise ValueError("Q4全年结果尚未人工验收，禁止生成正式提交")
    destination = root/f"outputs/submissions/result4-{variant}.xlsx"
    if destination.exists() or destination.is_symlink():
        raise FileExistsError(f"Q4提交文件已存在，禁止覆盖：{destination}")
    folder = resolve_submission_source(root, variant)
    outputs,metrics = load_outputs(folder)
    validate_submission_source(root,variant,outputs,metrics)
    if destination.resolve()!=root.resolve()/f"outputs/submissions/result4-{variant}.xlsx":
        raise ValueError("Q4提交路径指向了非指定目录")
    result = render_copy(root/f"data/raw/附件5/result4-{variant}.xlsx",destination,outputs,variant)
    from src.q3.search_runner import save_json
    check_protected(root,metrics)
    audit = root/"outputs/model_validation/final_export"
    audit.mkdir(parents=True,exist_ok=True)
    save_json(audit/f"q4_{variant}_submission_validation.json",result)
    return result
