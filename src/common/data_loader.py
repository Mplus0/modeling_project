"""只读加载官方 Excel；不改写表头、空白或单元格值。"""

from hashlib import sha256
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook


def discover_workbooks(raw_dir):
    root = Path(raw_dir).resolve()
    if not root.is_dir():
        raise FileNotFoundError(f"原始数据目录不存在: {root}")
    files = sorted(p for p in root.rglob("*") if p.is_file()
                   and p.suffix.lower() in {".xlsx", ".xlsm", ".xls"}
                   and not p.name.startswith("~$"))
    if not files:
        raise FileNotFoundError(f"未发现 Excel 附件: {root}")
    return files


def file_hash(path):
    with Path(path).open("rb") as stream:
        return sha256(stream.read()).hexdigest()


def load_sheets(path):
    if Path(path).suffix.lower() == ".xls":
        raise ValueError("不支持旧版 .xls；未转换文件。TODO: 需建模手确认")
    # 保留公式文本，避免把没有缓存值的公式误报为空白；从不调用 save。
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        for sheet in workbook:
            rows = list(sheet.iter_rows(values_only=True))
            headers = list(rows[0]) if rows else []
            # 用列位置作为内部索引，避免 pandas 自动重命名重复/空白表头。
            frame = pd.DataFrame(rows[1:], columns=range(len(headers)), dtype=object)
            yield {
                "sheet": sheet.title,
                "excel_rows": sheet.max_row,
                "excel_columns": sheet.max_column,
                "headers": headers,
                "frame": frame,
            }
    finally:
        workbook.close()
