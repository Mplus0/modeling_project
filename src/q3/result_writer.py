"""按团队确认口径填充官方result3副本，复用Q2样式块与紧急区间工具。"""

from datetime import datetime
from copy import copy
from pathlib import Path
from shutil import copy2
from tempfile import TemporaryDirectory

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from src.common.data_loader import file_hash
from src.q2.result_writer import capture_block, paste_block, clear_body, slot_intervals, emergency_intervals
from src.q3.annual import FORMAL_DATES
from src.q3.optimizer import VALIDATION_TOL
from src.q3.final import load_final, validate_final, read_confirmed_selection, compare_search_winner

SHEETS = ["计划购电量","调整购电量","充放电量","紧急购电量"]


def inspect_template(template):
    workbook = load_workbook(template)
    try:
        if workbook.sheetnames!=SHEETS:
            raise ValueError("Q3官方模板sheet名称或顺序不符")
        headers = []
        for name in SHEETS[:2]:
            sheet = workbook[name]
            if (sheet.max_row,sheet.max_column)!=(335,147):
                raise ValueError("Q3购电模板尺寸不符")
            header = [c.value for c in sheet[1]]
            if header[0]!="日期\\时间" or header[-2:]!=["全天购电量","全天购电费"]:
                raise ValueError("Q3购电模板标题不符")
            slot_intervals(header[1:145])
            if [pd.Timestamp(sheet.cell(i+2,1).value).date() for i in range(334)]!=list(FORMAL_DATES):
                raise ValueError("Q3购电模板日期不符")
            headers.append(header)
        if headers[0]!=headers[1]:
            raise ValueError("Q3计划及调整表时段标签不一致")
        if [c.value for c in workbook[SHEETS[2]][1]]!=["日期","时间段","充电量","放电量","时刻","储电量"]:
            raise ValueError("Q3充放电表标题不符")
        if [c.value for c in workbook[SHEETS[3]][1]]!=["日期","购电时间段","购电量"]:
            raise ValueError("Q3紧急表标题不符")
        return dict(sheets=SHEETS,dimensions={s.title:[s.max_row,s.max_column] for s in workbook},
                    mapping_confirmed=True,slot_mapping="官方slot原行序",charge="x_real+q_real",discharge="z_real",
                    planned="00:00 g0及计划费用",adjustment="三轮累计净变化sum(u-v)及sum(C*(1.5u-0.5v))",
                    expansion="仅输出副本扩展充放电全年块及逐日有效紧急区间",template_sha256=file_hash(template))
    finally:
        workbook.close()


def _display(value):
    return 0. if abs(float(value))<=VALIDATION_TOL else float(value)


def _populate(workbook, outputs):
    """内部映射函数只操作副本；公共入口先验收最终CSV与搜索一致性。"""
    planned,adjusted,battery,emergency = [workbook[s] for s in SHEETS]
    plans,actual,daily = [outputs[k] for k in ("plan_updates","actual_schedule","daily_metrics")]
    events = emergency_intervals(actual,slot_intervals([planned.cell(1,c).value for c in range(2,146)]))
    battery_block = capture_block(battery,2,7)
    emergency_blocks = [capture_block(emergency,r,r) for r in (2,3,4)]
    clear_body(battery)
    clear_body(emergency)
    for index,day in enumerate(FORMAL_DATES):
        label = str(day)
        p = plans[plans.date.eq(label)]
        original = p[p.update_time.eq("00:00")]
        changes = p[~p.update_time.eq("00:00")]
        # 同一slot的三轮增减累计为净变化，不能用最终有效合同量替代。
        net = (changes.increase_kwh-changes.decrease_kwh).groupby(changes.slot).sum().reindex(range(1,145),fill_value=0.)
        a = actual[actual.date.eq(label)]
        d = daily[daily.date.eq(label)].iloc[0]
        for col,(g,delta) in enumerate(zip(original.new_plan_kwh,net),2):
            planned.cell(index+2,col,_display(g))
            adjusted.cell(index+2,col,_display(delta))
        planned.cell(index+2,146,_display(original.new_plan_kwh.sum()))
        planned.cell(index+2,147,_display(d.plan_cost_00_yuan))
        adjusted.cell(index+2,146,_display(net.sum()))
        adjusted.cell(index+2,147,_display(changes.adjustment_cost_yuan.sum()))
        first = 2+index*6
        paste_block(battery,battery_block,first)
        battery.cell(first,1,datetime.combine(day,datetime.min.time()))
        for block in range(6):
            part = a.iloc[block*24:(block+1)*24]
            battery.cell(first+block,3,_display((part.x_real_kwh+part.q_real_kwh).sum()))
            battery.cell(first+block,4,_display(part.z_real_kwh.sum()))
        battery.cell(first,6,_display(a.soc_real_start_kwh.iloc[0]))
        battery.cell(first+1,6,_display(a.soc_real_end_kwh.iloc[-1]))
    for i,event in enumerate(events):
        first_day = i==0 or events[i-1]["date"]!=event["date"]
        last_day = i==len(events)-1 or events[i+1]["date"]!=event["date"]
        paste_block(emergency,emergency_blocks[0 if first_day else (2 if last_day else 1)],i+2)
        emergency.cell(i+2,1,datetime.fromisoformat(event["date"]))
        emergency.cell(i+2,2,event["interval"])
        emergency.cell(i+2,3,_display(event["emergency_purchase_kwh"]))
    return events


def verify_workbook(template, destination, outputs):
    def style(cell):
        # 保存后样式索引可能重排，比较实际格式属性而非内部索引编号。
        return (copy(cell.font),copy(cell.fill),copy(cell.border),copy(cell.alignment),copy(cell.protection),
                cell.number_format,cell.quotePrefix,cell.pivotButton)
    original,filled = load_workbook(template),load_workbook(destination)
    try:
        if filled.sheetnames!=original.sheetnames:
            raise ValueError("Q3提交sheet结构改变")
        for name in SHEETS:
            a,b = original[name],filled[name]
            if [c.value for c in a[1]]!=[c.value for c in b[1]] or a.max_column!=b.max_column:
                raise ValueError("Q3提交标题或列顺序改变")
            for col in range(1,a.max_column+1):
                if style(a.cell(1,col))!=style(b.cell(1,col)):
                    raise ValueError("Q3提交标题格式改变")
        if filled["充放电量"].max_row!=2005:
            raise ValueError("Q3充放电表不是334个六行日块")
        expected = load_workbook(template)
        try:
            events = _populate(expected,outputs)
            for name in SHEETS:
                a,b = expected[name],filled[name]
                if (a.max_row,a.max_column)!=(b.max_row,b.max_column) or str(a.merged_cells)!=str(b.merged_cells):
                    raise ValueError("Q3提交扩展结构与模板规则不一致")
                for rows in zip(a.iter_rows(),b.iter_rows()):
                    for x,y in zip(*rows):
                        if isinstance(x.value,(int,float)):
                            if not isinstance(y.value,(int,float)) or not np.isfinite(y.value) or abs(x.value-y.value)>VALIDATION_TOL:
                                raise ValueError(f"Q3提交数值不符:{name}!{x.coordinate}")
                        elif x.value!=y.value:
                            raise ValueError(f"Q3提交标签不符:{name}!{x.coordinate}")
                        if x._style!=y._style and style(x)!=style(y):
                            raise ValueError(f"Q3提交样式不符:{name}!{x.coordinate}")
                for r in a.row_dimensions:
                    if a.row_dimensions[r].height!=b.row_dimensions[r].height:
                        raise ValueError("Q3提交行高不符")
            return dict(passed=True,days=334,slots=48096,emergency_intervals=len(events),sheet_names=SHEETS)
        finally:
            expected.close()
    finally:
        original.close()
        filled.close()


def write_result3(root):
    directory = root/"outputs/q3/final"
    outputs,metrics = load_final(directory)
    validate_final(outputs,metrics)
    _,_,winner_daily,winner_metrics = read_confirmed_selection(root)
    if not compare_search_winner(outputs["daily_metrics"],metrics,winner_daily,winner_metrics)["passed"]:
        raise ValueError("Q3最终结果与搜索winner不一致，禁止生成submission")
    template,destination = root/"data/raw/附件5/result3.xlsx",root/"outputs/submissions/result3.xlsx"
    inspect_template(template)
    if destination.resolve()!=root.resolve()/"outputs/submissions/result3.xlsx":
        raise ValueError("Q3提交路径不得指向raw或其他文件")
    before = file_hash(template)
    destination.parent.mkdir(parents=True,exist_ok=True)
    with TemporaryDirectory(prefix=".q3-result-",dir=destination.parent) as temporary:
        staged = Path(temporary)/"result3.xlsx"
        copy2(template,staged)
        workbook = load_workbook(staged)
        try:
            _populate(workbook,outputs)
            workbook.save(staged)
        finally:
            workbook.close()
        checked = verify_workbook(template,staged,outputs)
        if file_hash(template)!=before:
            raise RuntimeError("Q3原始模板SHA-256变化")
        staged.replace(destination)
    return dict(checked,status="generated",template_sha256=before,submission_sha256=file_hash(destination))
